import openpyxl
import pandas as pd
import sys

backend_dir = r"c:\Users\hp\Desktop\Audit\backend"
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.loaders import load_file

po_path = r"C:\Users\hp\Documents\SQL Server Management Studio\P2P New Q2 Data\Purchase Order Report-ITL.xls"
ref_path = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"

with open(po_path, 'rb') as f:
    df_po = load_file(f, "Purchase Order Report-ITL.xls")

wb_ref = openpyxl.load_workbook(ref_path, data_only=True)
ws_ref = wb_ref['Purchase Order Report-ITL']

print("=== FIRST 10 ROWS COMPARISON ===")
for r in range(5, 15):
    ref_po_no = ws_ref.cell(row=r, column=2).value
    ref_item = ws_ref.cell(row=r, column=15).value
    ref_qty = ws_ref.cell(row=r, column=18).value
    ref_price = ws_ref.cell(row=r, column=26).value
    ref_line_tot = ws_ref.cell(row=r, column=33).value
    
    # Matching row in df_po
    po_row = df_po.iloc[r-5]
    print(f"Row {r:2d}: Ref PO={ref_po_no}, Item={ref_item}, Qty={ref_qty}, Price={ref_price}, LineTot={ref_line_tot}")
    print(f"        PO  PO={po_row['PO No']}, Item={po_row['Item Code']}, Qty={po_row['PO Qty']}, Price={po_row['PO Price']}, LineTot={po_row['Line Total']}")

# Total sum of Line Total in ref sheet
ref_line_tots = [ws_ref.cell(row=r, column=33).value for r in range(5, ws_ref.max_row + 1) if ws_ref.cell(row=r, column=33).value is not None]
ref_sum = sum(ref_line_tots)
print(f"\nSum of Line Total in Ref Sheet: {ref_sum} ({ref_sum / 10000000:.2f} Cr)")

po_line_tots = pd.to_numeric(df_po['Line Total'], errors='coerce').fillna(0)
print(f"Sum of Line Total in df_po raw: {po_line_tots.sum()} ({po_line_tots.sum() / 10000000:.2f} Cr)")
