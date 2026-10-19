import openpyxl

filepath = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_f = openpyxl.load_workbook(filepath, data_only=False)
wb_v = openpyxl.load_workbook(filepath, data_only=True)

ws_f = wb_f['Purchase Order Report-ITL']
ws_v = wb_v['Purchase Order Report-ITL']

kpis = []
for col in range(1, 13):
    title = ws_v.cell(row=1, column=col).value
    f_val = ws_f.cell(row=2, column=col).value
    v_val = ws_v.cell(row=2, column=col).value
    
    kpis.append({
        "index": col,
        "col_letter": openpyxl.utils.get_column_letter(col),
        "title": title,
        "formula": str(f_val),
        "evaluated_value": v_val
    })

print("=== 12 KPIS IN REFERENCE FILE ===")
for k in kpis:
    print(f"KPI {k['index']:2d} [{k['col_letter']}]: '{k['title']}' | Formula: {k['formula']} | Value: {k['evaluated_value']}")

print("\nCheck if there are any other KPIs or tables on Sheet3 or HOLIDAY-2026:")
for sname in wb_f.sheetnames:
    ws = wb_f[sname]
    print(f"Sheet '{sname}' max row: {ws.max_row}, max col: {ws.max_column}")
