import openpyxl
import pandas as pd
import numpy as np

filepath = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_v = openpyxl.load_workbook(filepath, data_only=True)
ws_v = wb_v['Purchase Order Report-ITL']

ref_aks = []
ags = []
ahs = []

for r in range(5, ws_v.max_row + 1):
    ref_ak = ws_v.cell(row=r, column=37).value # Col AK
    ag = ws_v.cell(row=r, column=33).value # Col AG
    ah = ws_v.cell(row=r, column=34).value # Col AH
    ref_aks.append(ref_ak)
    ags.append(ag)
    ahs.append(ah)

df = pd.DataFrame({"ref_ak": ref_aks, "AG": ags, "AH": ahs})
df['AG'] = pd.to_numeric(df['AG'], errors='coerce')
df['AH'] = pd.to_numeric(df['AH'], errors='coerce')

# What if we round(abs(AG - AH) / AG, 6) or round(diff, 4)?
diff = df['AG'] - df['AH']
pct = (df['AH'] - df['AG']) / df['AG']

# Test round(pct, 6) <= 0.05
pct_round6 = pct.round(6)
df['calc_r6'] = np.where(df['AH'].isna(), "Regular", np.where(diff >= 0, "Regular", np.where(pct_round6 <= 0.05, "Tolerable", "Exception")))

# Test round(pct, 10) <= 0.05
pct_round10 = pct.round(10)
df['calc_r10'] = np.where(df['AH'].isna(), "Regular", np.where(diff >= 0, "Regular", np.where(pct_round10 <= 0.05, "Tolerable", "Exception")))

# Test exact match to ref_ak
print("ref_ak Exception count:", (df['ref_ak'] == 'Exception').sum())
print("calc_r6 Exception count:", (df['calc_r6'] == 'Exception').sum(), "| Total Matches to ref_ak:", (df['ref_ak'] == df['calc_r6']).sum(), "/", len(df))
print("calc_r10 Exception count:", (df['calc_r10'] == 'Exception').sum(), "| Total Matches to ref_ak:", (df['ref_ak'] == df['calc_r10']).sum(), "/", len(df))

# Test numpy isclose or small epsilon
pct_eps = pct - 1e-12
df['calc_eps'] = np.where(df['AH'].isna(), "Regular", np.where(diff >= 0, "Regular", np.where(pct_eps <= 0.05, "Tolerable", "Exception")))
print("calc_eps Exception count:", (df['calc_eps'] == 'Exception').sum(), "| Total Matches to ref_ak:", (df['ref_ak'] == df['calc_eps']).sum(), "/", len(df))
