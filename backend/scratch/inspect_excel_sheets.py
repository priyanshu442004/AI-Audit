import pandas as pd
import openpyxl

file_path = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"

wb = openpyxl.load_workbook(file_path, data_only=True)
print("SHEET NAMES IN Purchase Status Report.xlsx:")
for sheet in wb.sheetnames:
    print(f" - {sheet}")

for sheet_name in wb.sheetnames:
    print(f"\n==========================================")
    print(f"SHEET: {sheet_name}")
    print(f"==========================================")
    df = pd.read_excel(file_path, sheet_name=sheet_name)
    print("Shape:", df.shape)
    print("Columns:", list(df.columns[:15]))
    print("Head (3 rows):")
    print(df.head(3))
