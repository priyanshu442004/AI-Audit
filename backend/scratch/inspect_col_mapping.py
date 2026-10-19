import openpyxl
import pandas as pd
import sys

backend_dir = r"c:\Users\hp\Desktop\Audit\backend"
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.loaders import load_file

po_path = r"C:\Users\hp\Documents\SQL Server Management Studio\P2P New Q2 Data\Purchase Order Report-ITL.xls"
ref_path = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"

with open(po_path, 'rb') as f:
    df_po = load_file(f, "Purchase Order Report-ITL.xls")

wb_ref = openpyxl.load_workbook(ref_path, data_only=True)
ws_ref = wb_ref['Purchase Order Report-ITL']

print(f"PO File columns ({len(df_po.columns)}):")
for i, c in enumerate(df_po.columns):
    print(f"  Col {i:2d}: '{c}'")

print(f"\nRef File Header (Row 4) & Row 5:")
for c in range(1, 39):
    c_let = openpyxl.utils.get_column_letter(c)
    h_name = ws_ref.cell(row=4, column=c).value
    r5_val = ws_ref.cell(row=5, column=c).value
    print(f"  Col {c:2d} ({c_let}): '{h_name}' | R5: {r5_val}")
