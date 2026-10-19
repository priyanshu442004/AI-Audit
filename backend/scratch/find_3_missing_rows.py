import openpyxl
import pandas as pd
import numpy as np
import sys

backend_dir = r"c:\Users\hp\Desktop\Audit\backend"
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from scratch.test_po_status_logic import df_res, df_po

ref_path = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_ref = openpyxl.load_workbook(ref_path, data_only=True)
ws_ref = wb_ref['Purchase Order Report-ITL']

mismatches = []
for r in range(5, ws_ref.max_row + 1):
    ref_ak = ws_ref.cell(row=r, column=37).value # Col AK: Line Total Difference
    ref_ag = ws_ref.cell(row=r, column=33).value # Col AG: Line Total
    ref_ah = ws_ref.cell(row=r, column=34).value # Col AH: GRN Line Total
    ref_ai = ws_ref.cell(row=r, column=35).value # Col AI: Excess Price
    
    py_row = df_res.iloc[r-5]
    py_ak = py_row['Line Total Difference']
    py_ag = py_row['Line Total']
    py_ah = py_row['GRN Line Total']
    py_ai = py_row['Excess Price']
    
    if ref_ak != py_ak:
        mismatches.append({
            "row": r,
            "po_no": ws_ref.cell(row=r, column=2).value,
            "item_code": ws_ref.cell(row=r, column=15).value,
            "ref_ak": ref_ak,
            "py_ak": py_ak,
            "ref_ag": ref_ag,
            "py_ag": py_ag,
            "ref_ah": ref_ah,
            "py_ah": py_ah,
            "ref_ai": ref_ai,
            "py_ai": py_ai
        })

print(f"Total Mismatches between ref sheet Col AK and Python Col AK: {len(mismatches)}")
for m in mismatches:
    print(m)
