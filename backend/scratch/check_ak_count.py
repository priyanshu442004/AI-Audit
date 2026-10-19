import openpyxl
import pandas as pd
import numpy as np

filepath = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_val = openpyxl.load_workbook(filepath, data_only=True)
ws_v = wb_val['Purchase Order Report-ITL']

exc_count = 0
for r in range(5, ws_v.max_row + 1):
    ak = ws_v.cell(row=r, column=37).value
    if ak == 'Exception':
        exc_count += 1

print(f"Total 'Exception' in Reference Sheet Col AK: {exc_count}")
