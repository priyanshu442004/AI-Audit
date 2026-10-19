import openpyxl
import pandas as pd

filepath = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_val = openpyxl.load_workbook(filepath, data_only=True)
ws_v = wb_val['Purchase Order Report-ITL']

ref_holidays = []
ref_sundays = []

for r in range(5, ws_v.max_row + 1):
    po_no = ws_v.cell(row=r, column=2).value # B
    is_num = isinstance(po_no, (int, float)) and not isinstance(po_no, bool)
    post_dt = ws_v.cell(row=r, column=5).value # E
    exc = ws_v.cell(row=r, column=21).value # U
    
    if exc == 'Sunday':
        ref_sundays.append((r, po_no, is_num, post_dt, exc))
    elif exc not in ['Sunday', 'Regular', None]:
        ref_holidays.append((r, po_no, is_num, post_dt, exc))

print(f"Total Sunday rows in Ref Col U: {len(ref_sundays)}")
print(f"Total Holiday rows in Ref Col U: {len(ref_holidays)}")

unique_sundays_all = len(set(x[1] for x in ref_sundays if x[1] not in [None, '', 0, '0']))
unique_sundays_num_only = len(set(x[1] for x in ref_sundays if x[2]))

unique_holidays_all = len(set(x[1] for x in ref_holidays if x[1] not in [None, '', 0, '0']))
unique_holidays_num_only = len(set(x[1] for x in ref_holidays if x[2]))

print(f"Unique Sundays (All): {unique_sundays_all} | (Numeric PO only): {unique_sundays_num_only}")
print(f"Unique Holidays (All): {unique_holidays_all} | (Numeric PO only): {unique_holidays_num_only}")

print("\nSample Sunday rows in ref:")
for s in ref_sundays[:5]:
    print(" ", s)

print("\nSample Holiday rows in ref:")
for h in ref_holidays[:5]:
    print(" ", h)
