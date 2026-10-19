import openpyxl
import pandas as pd
import numpy as np

filepath = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_f = openpyxl.load_workbook(filepath, data_only=False)
wb_v = openpyxl.load_workbook(filepath, data_only=True)

ws_f = wb_f['Purchase Order Report-ITL']
ws_v = wb_v['Purchase Order Report-ITL']

mismatches = []
for r in range(5, ws_v.max_row + 1):
    ref_ak = ws_v.cell(row=r, column=37).value # Col AK: Line Total Difference
    ref_ag = ws_v.cell(row=r, column=33).value # Col AG: Line Total
    ref_ah = ws_v.cell(row=r, column=34).value # Col AH: GRN Line Total
    ref_ai = ws_v.cell(row=r, column=35).value # Col AI: Excess Price
    
    # Let's compute exact formula:
    # IF ISNA(AH) -> Regular
    # IF AG - AH >= 0 -> Regular
    # IF ABS(AG - AH) / AG <= 0.05 -> Tolerable
    # ELSE -> Exception
    
    if ref_ah is None or str(ref_ah).startswith('#') or pd.isna(ref_ah):
        calc_ak = "Regular"
    elif ref_ag - ref_ah >= 0:
        calc_ak = "Regular"
    else:
        diff_pct = abs(ref_ag - ref_ah) / ref_ag
        if diff_pct <= 0.05:
            calc_ak = "Tolerable"
        else:
            calc_ak = "Exception"
            
    if calc_ak != ref_ak:
        mismatches.append((r, ref_ag, ref_ah, ref_ai, ref_ak, calc_ak, abs(ref_ag - ref_ah) / ref_ag))

print(f"Total Mismatches when evaluating directly on Excel values: {len(mismatches)}")
for m in mismatches:
    print("  Row:", m[0], "| AG (LineTot):", m[1], "| AH (GRN LineTot):", m[2], "| AI (ExcPrice):", m[3], "| Ref AK:", m[4], "| Calc AK:", m[5], "| Pct:", m[6])
