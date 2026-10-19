import openpyxl
import pandas as pd
import sys

backend_dir = r"c:\Users\hp\Desktop\Audit\backend"
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from scratch.test_po_status_logic import df_res

ref_path = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_ref = openpyxl.load_workbook(ref_path, data_only=True)
ws_ref = wb_ref['Purchase Order Report-ITL']

diff_rows = []
for r in range(5, ws_ref.max_row + 1):
    ref_key = ws_ref.cell(row=r, column=1).value
    ref_line_diff = ws_ref.cell(row=r, column=37).value # Col AK
    ref_exc_price = ws_ref.cell(row=r, column=35).value # Col AI
    
    py_row = df_res.iloc[r-5]
    py_line_diff = py_row['Line Total Difference']
    py_exc_price = py_row['Excess Price']
    
    if ref_line_diff != py_line_diff:
        ref_po_line_tot = ws_ref.cell(row=r, column=33).value
        ref_grn_line_tot = ws_ref.cell(row=r, column=34).value
        diff_rows.append({
            "row": r,
            "key": ref_key,
            "ref_line_diff": ref_line_diff,
            "py_line_diff": py_line_diff,
            "ref_po_line_tot": ref_po_line_tot,
            "ref_grn_line_tot": ref_grn_line_tot,
            "py_po_line_tot": py_row['Line Total'],
            "py_grn_line_tot": py_row['GRN Line Total'],
            "ref_exc_price": ref_exc_price,
            "py_exc_price": py_exc_price
        })

print(f"Total Discrepancy Rows in Line Total Difference: {len(diff_rows)}")
for d in diff_rows:
    print(d)
