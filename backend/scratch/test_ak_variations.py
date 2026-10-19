import openpyxl
import pandas as pd
import numpy as np
from decimal import Decimal

filepath = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"
wb_v = openpyxl.load_workbook(filepath, data_only=True)
ws_v = wb_v['Purchase Order Report-ITL']

ref_aks = []
ags = []
ahs = []

for r in range(5, ws_v.max_row + 1):
    ref_ak = ws_v.cell(row=r, column=37).value # Col AK
    ag = ws_v.cell(row=r, column=33).value # Col AG
    ah = ws_v.cell(row=r, column=34).value # Col AH
    ref_aks.append(ref_ak)
    ags.append(ag)
    ahs.append(ah)

df = pd.DataFrame({"ref_ak": ref_aks, "AG": ags, "AH": ahs})

df['AG'] = pd.to_numeric(df['AG'], errors='coerce')
df['AH'] = pd.to_numeric(df['AH'], errors='coerce')

# Condition 1: Strict <= 0.05
diff_pct1 = (df['AH'] - df['AG']) / df['AG']
df['calc1'] = np.where(df['AH'].isna(), "Regular", np.where(df['AG'] - df['AH'] >= 0, "Regular", np.where(diff_pct1 <= 0.05, "Tolerable", "Exception")))

# Condition 2: Decimal exact <= 0.05
calcs3 = []
for idx, row in df.iterrows():
    ag, ah = row['AG'], row['AH']
    if pd.isna(ah) or ah is None:
        calcs3.append("Regular")
    elif ag - ah >= 0:
        calcs3.append("Regular")
    else:
        dec_ag = Decimal(str(ag))
        dec_ah = Decimal(str(ah))
        pct = (dec_ah - dec_ag) / dec_ag
        if pct <= Decimal('0.05'):
            calcs3.append("Tolerable")
        else:
            calcs3.append("Exception")

df['calc3'] = calcs3

print("Ref count Exception:", (df['ref_ak'] == 'Exception').sum())
print("Calc1 count Exception (strict float <= 0.05):", (df['calc1'] == 'Exception').sum())
print("Calc3 count Exception (Decimal exact <= 0.05):", (df['calc3'] == 'Exception').sum())

print("Differences for Calc1:")
print(df[df['ref_ak'] != df['calc1']][['ref_ak', 'calc1', 'AG', 'AH']])

print("Differences for Calc3:")
print(df[df['ref_ak'] != df['calc3']][['ref_ak', 'calc3', 'AG', 'AH']])
