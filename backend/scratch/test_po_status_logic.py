import os
import sys
import pandas as pd
import numpy as np
import openpyxl

backend_dir = r"c:\Users\hp\Desktop\Audit\backend"
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.loaders import load_file

po_path = r"C:\Users\hp\Documents\SQL Server Management Studio\P2P New Q2 Data\Purchase Order Report-ITL.xls"
grpo_path = r"C:\Users\hp\Documents\SQL Server Management Studio\P2P New Q2 Data\new GRPO Report-ITL1.xls"
holiday_path = r"C:\Users\hp\Documents\SQL Server Management Studio\P2P New Q2 Data\HOLIDAY-2026.xlsx"
ref_path = r"C:\Users\hp\Downloads\Purchase Status Report.xlsx"

# Helper to safely clean numeric series with commas
def clean_num(series):
    if series is None:
        return 0
    return pd.to_numeric(series.astype(str).str.replace(',', '').str.strip(), errors='coerce')

def safe_parse_date(series):
    return pd.to_datetime(series, dayfirst=True, errors='coerce')

# Load files
with open(po_path, 'rb') as f:
    df_po = load_file(f, "Purchase Order Report-ITL.xls")

with open(grpo_path, 'rb') as f:
    df_grpo = load_file(f, "new GRPO Report-ITL1.xls")

df_holiday = pd.read_excel(holiday_path)

# Prepare GRPO lookups
grpo_po_no_col = 'PO Number' if 'PO Number' in df_grpo.columns else 'PO No'
grpo_no_col = 'GRPO No' if 'GRPO No' in df_grpo.columns else 'GRN No'
grpo_item_col = 'Item Code' if 'Item Code' in df_grpo.columns else 'Item No.'
grpo_qty_col = 'GRPO Qty' if 'GRPO Qty' in df_grpo.columns else 'Quantity'
grpo_price_col = 'GRPO Price' if 'GRPO Price' in df_grpo.columns else 'Unit Price'
grpo_line_total_col = 'Line Total' if 'Line Total' in df_grpo.columns else 'Row Total'

# Clean numeric types in GRPO
df_grpo[grpo_qty_col] = clean_num(df_grpo[grpo_qty_col])
df_grpo[grpo_price_col] = clean_num(df_grpo[grpo_price_col])
df_grpo[grpo_line_total_col] = clean_num(df_grpo[grpo_line_total_col])

# Create Key in GRPO: str(PO Number) + "-" + str(Item Code)
df_grpo['_key'] = df_grpo[grpo_po_no_col].astype(str).str.strip() + "-" + df_grpo[grpo_item_col].astype(str).str.strip()
df_grpo_dedup = df_grpo.drop_duplicates(subset=['_key'], keep='first')

grpo_key_qty = df_grpo_dedup.set_index('_key')[grpo_qty_col].to_dict()
grpo_key_price = df_grpo_dedup.set_index('_key')[grpo_price_col].to_dict()
grpo_key_line_total = df_grpo_dedup.set_index('_key')[grpo_line_total_col].to_dict()

# First GRPO No per PO No
df_grpo_po_first = df_grpo.dropna(subset=[grpo_po_no_col]).drop_duplicates(subset=[grpo_po_no_col], keep='first')
grpo_po_to_grn_no = df_grpo_po_first.set_index(grpo_po_no_col)[grpo_no_col].to_dict()

# Holiday dates set & map
holiday_map = {}
for idx, row in df_holiday.iterrows():
    d_val = row.iloc[-1]
    h_name = row.iloc[1]
    if pd.notna(d_val):
        d_obj = pd.to_datetime(d_val).date()
        holiday_map[d_obj] = str(h_name).strip()

# Process PO dataframe
po_no_col = 'PO No'
item_code_col = 'Item Code'
posting_date_col = 'Posting Date'
canceled_status_col = 'Canceled Status'
doc_status_col = 'Document Status'
po_qty_col = 'PO Qty'
open_qty_col = 'Open Qty'
po_price_col = 'PO Price'
line_total_col = 'Line Total'

df_res = pd.DataFrame()
df_res['PO No'] = df_po[po_no_col].astype(str).str.strip()
df_res['Item Code'] = df_po[item_code_col].astype(str).str.strip()
df_res['Key'] = df_res['PO No'] + "-" + df_res['Item Code']
df_res['GRN No'] = df_res['PO No'].map(grpo_po_to_grn_no)

df_res['PO Series'] = df_po.get('PO Series', '')
df_res['Posting Date'] = safe_parse_date(df_po[posting_date_col])
df_res['Delivery Date'] = safe_parse_date(df_po.get('Delivery Date'))
df_res['Document Date'] = safe_parse_date(df_po.get('Document Date'))
df_res['Branch'] = df_po.get('Branch', '')
df_res['Document Status'] = df_po[doc_status_col].astype(str).str.strip()
df_res['Canceled Status'] = df_po[canceled_status_col].astype(str).str.strip()
df_res['Vendor Group'] = df_po.get('Vendor Group', '')
df_res['Vendor Code'] = df_po.get('Vendor Code', '')
df_res['Vendor Name'] = df_po.get('Vendor Name', '')
df_res['Item Group'] = df_po.get('Item Group', '')
df_res['Item Description'] = df_po.get('Item Description', '')
df_res['UOM'] = df_po.get('UOM', '')

df_res['PO Qty'] = clean_num(df_po[po_qty_col]).fillna(0)
df_res['Open Qty'] = clean_num(df_po[open_qty_col]).fillna(0)

# Matched GRN fields
df_res['GRN Qty'] = df_res['Key'].map(grpo_key_qty)

# PO Monitoring Status
def calc_po_mon_status(row):
    if row['Canceled Status'] == 'Yes':
        return 'PO Canceled'
    if pd.isna(row['GRN Qty']):
        return 'Item Not Received'
    if row['PO Qty'] == row['Open Qty'] and row['GRN Qty'] > 0:
        return 'Open Qty Error'
    if row['Document Status'] == 'Closed':
        return 'PO Closed'
    if row['Open Qty'] == 0:
        return 'Fully Received'
    if row['Document Status'] == 'Open' and row['PO Qty'] > row['Open Qty'] and row['Open Qty'] > 0:
        return 'Partially Received'
    return 'Check'

df_res['PO Monitoring Status'] = df_res.apply(calc_po_mon_status, axis=1)

# Holiday & Sunday Exception
def calc_holiday_exc(dt):
    if pd.isna(dt):
        return 'Regular'
    d = dt.date()
    if dt.weekday() == 6: # Sunday
        return 'Sunday'
    if d in holiday_map:
        return f"Holiday - {holiday_map[d]}"
    return 'Regular'

df_res['Holiday & Sunday Exception'] = df_res['Posting Date'].apply(calc_holiday_exc)

# Excess QTY: IFERROR(IF(PO_Qty - GRN_Qty < 0, ABS(PO_Qty - GRN_Qty), 0), 0)
df_res['Excess QTY'] = np.where(
    df_res['GRN Qty'].isna(),
    0,
    np.where(df_res['PO Qty'] - df_res['GRN Qty'] < 0, (df_res['PO Qty'] - df_res['GRN Qty']).abs(), 0)
)

# Excess Qty Variation %: IF(W>0, W/R, 0)
df_res['Excess Qty Variation %'] = np.where(df_res['Excess QTY'] > 0, df_res['Excess QTY'] / df_res['PO Qty'], 0)

# QTY Difference: IF(ISNA(V),"Regular",IF(R-V>=0,"Regular",IF(ABS(R-V)/R<=5%,"Tolerable","Exception")))
def calc_qty_diff(row):
    if pd.isna(row['GRN Qty']):
        return 'Regular'
    diff = row['PO Qty'] - row['GRN Qty']
    if diff >= 0:
        return 'Regular'
    if row['PO Qty'] > 0 and (abs(diff) / row['PO Qty']) <= 0.05 + 1e-9:
        return 'Tolerable'
    return 'Exception'

df_res['QTY Difference'] = df_res.apply(calc_qty_diff, axis=1)

# Prices
df_res['PO Price'] = clean_num(df_po[po_price_col]).fillna(0)
df_res['GRN Price'] = df_res['Key'].map(grpo_key_price)

df_res['Excess Rate'] = np.where(
    df_res['GRN Price'].isna(),
    0,
    np.where(df_res['PO Price'] - df_res['GRN Price'] < 0, (df_res['PO Price'] - df_res['GRN Price']).abs(), 0)
)
df_res['Excess Rate Variation %'] = np.where(df_res['Excess Rate'] > 0, df_res['Excess Rate'] / df_res['PO Price'], 0)

def calc_rate_diff(row):
    if pd.isna(row['GRN Price']):
        return 'Regular'
    diff = row['PO Price'] - row['GRN Price']
    if diff >= 0:
        return 'Regular'
    if row['PO Price'] > 0 and (abs(diff) / row['PO Price']) <= 0.05 + 1e-9:
        return 'Tolerable'
    return 'Exception'

df_res['Rate Difference'] = df_res.apply(calc_rate_diff, axis=1)

df_res['Document Currency'] = df_po.get('Document Currency', 'INR')
df_res['Document Rate'] = clean_num(df_po.get('Document Rate', 1)).fillna(1)
df_res['Line Total'] = clean_num(df_po[line_total_col]).fillna(df_res['PO Qty'] * df_res['PO Price'])
df_res['GRN Line Total'] = df_res['Key'].map(grpo_key_line_total)

df_res['Excess Price'] = np.where(
    df_res['GRN Line Total'].isna(),
    0,
    np.where(df_res['Line Total'] - df_res['GRN Line Total'] < 0, (df_res['Line Total'] - df_res['GRN Line Total']).abs(), 0)
)
df_res['Excess Price Variation %'] = np.where(df_res['Excess Price'] > 0, df_res['Excess Price'] / df_res['Line Total'], 0)

def calc_line_diff(row):
    if pd.isna(row['GRN Line Total']):
        return 'Regular'
    diff = row['Line Total'] - row['GRN Line Total']
    if diff >= 0:
        return 'Regular'
    if row['Line Total'] > 0 and (abs(diff) / row['Line Total']) <= 0.05 + 1e-9:
        return 'Tolerable'
    return 'Exception'

df_res['Line Total Difference'] = df_res.apply(calc_line_diff, axis=1)
df_res['Document Total'] = clean_num(df_po.get('Document Total', 0)).fillna(0)

# Read reference workbook evaluated values for comparison
wb_ref = openpyxl.load_workbook(ref_path, data_only=True)
ws_ref = wb_ref['Purchase Order Report-ITL']

ref_kpis = {}
for c in range(1, 13):
    title = ws_ref.cell(row=1, column=c).value
    val = ws_ref.cell(row=2, column=c).value
    ref_kpis[title] = val

print("\n=== CALCULATED KPIS VS REFERENCE KPIS ===")

valid_po_mask = (df_res['PO No'] != '') & (df_res['PO No'] != '0') & (df_res['PO No'].notna())

unique_po_count = df_res[valid_po_mask]['PO No'].nunique()
total_po_val_cr = round(df_res['Line Total'].sum() / 10000000, 2)

canceled_po_count = df_res[valid_po_mask & (df_res['Canceled Status'] == 'Yes')]['PO No'].nunique()

# Holiday Exception formula: count unique POs that fall on a Holiday (excluding Sunday and Regular)
holiday_po_mask = valid_po_mask & (~df_res['Holiday & Sunday Exception'].isin(['Sunday', 'Regular']))
holiday_po_count = df_res[holiday_po_mask]['PO No'].nunique()

# Sunday Exception formula: count unique POs that fall on Sunday
sunday_po_mask = valid_po_mask & (df_res['Holiday & Sunday Exception'] == 'Sunday')
sunday_po_count = df_res[sunday_po_mask]['PO No'].nunique()

cnt_excess_qty = (df_res['QTY Difference'] == 'Exception').sum()
sum_excess_qty_lakh = round(df_res[df_res['QTY Difference'] == 'Exception']['Excess QTY'].sum() / 100000, 2)

cnt_excess_rate = (df_res['Rate Difference'] == 'Exception').sum()

exc_rate_rows = df_res[df_res['Rate Difference'] == 'Exception']
sum_excess_rate = round((exc_rate_rows['Excess Rate'] * exc_rate_rows['Document Rate']).sum(), 2)

cnt_excess_price = (df_res['Line Total Difference'] == 'Exception').sum()
sum_excess_price_lakh = round(df_res[df_res['Line Total Difference'] == 'Exception']['Excess Price'].sum() / 100000, 2)

kpi_results = {
    "Foreign Vendor": "FILTER",
    "Unique PO Raise": unique_po_count,
    "Total PO Value": f"{total_po_val_cr:.2f} Cr",
    "Cancled PO Count": canceled_po_count,
    "Holiday Exception": holiday_po_count,
    "Sunday Exception": sunday_po_count,
    "Count of Excess QTY": int(cnt_excess_qty),
    "Sum of Excess QTY": f"{sum_excess_qty_lakh:.2f} Lakh",
    "Count of Excess Rate": int(cnt_excess_rate),
    "Sum of Excess Rate": float(sum_excess_rate),
    "Count of Excess Price": int(cnt_excess_price),
    "Sum of Excess Price": f"{sum_excess_price_lakh:.2f} Lakh",
}

print(f"{'KPI Title':<25s} | {'Calculated':<15s} | {'Reference':<15s} | {'Match?':<10s}")
print("-" * 75)
all_matched = True
for k, ref_v in ref_kpis.items():
    calc_v = kpi_results.get(k, "N/A")
    match = (str(calc_v) == str(ref_v))
    if not match:
        all_matched = False
    print(f"{str(k):<25s} | {str(calc_v):<15s} | {str(ref_v):<15s} | {'[MATCH]' if match else '[DIFF]':<10s}")

print(f"\nALL 12 KPIS PERFECTLY MATCHED: {all_matched}")
